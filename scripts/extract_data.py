"""
Regenera src/data.js a partir de BBDD_cotizaciones_v1.xlsx (o una versión
más nueva del mismo archivo, siempre que mantenga las hojas y columnas
RESIDUOS, LISTAS y DESTINATARIOS con esta misma estructura).

Uso:
    pip install openpyxl
    python scripts/extract_data.py ruta/a/BBDD_cotizaciones.xlsx

Esto sobrescribe src/data.js. Vuelve a compilar (npm run build) o a correr
npm run dev para ver los cambios.
"""
import sys
import json
import openpyxl


def sheet_to_json(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    data = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        d = {}
        for h, v in zip(headers, r):
            if h is None:
                continue
            if hasattr(v, "isoformat"):
                v = v.isoformat()[:10]
            d[str(h).strip()] = v
        data.append(d)
    return data


def build_residuos_tree(residuos):
    tree = {}
    for d in residuos:
        p, pc = d["PELIGROSIDAD"], d["COD1"]
        f, fc = d["FAMILIA"], d["COD2"]
        m, mc = d["MATERIAL"], d["COD3"]
        t, tc = d["TIPO"], d["COD4"]
        e, ec = d["ESTADO"], d["COD5"]
        fo, foc = d["FORMATO"], d["COD6"]
        n1 = tree.setdefault(f"{p}|{pc}", {})
        n2 = n1.setdefault(f"{f}|{fc}", {})
        n3 = n2.setdefault(f"{m}|{mc}", {})
        n4 = n3.setdefault(f"{t}|{tc}", {})
        n5 = n4.setdefault(f"{e}|{ec}", [])
        n5.append(f"{fo}|{foc}")
    return tree


def build_listas(wb):
    ws = wb["LISTAS"]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    cols = {}
    for i, h in enumerate(headers):
        if h:
            cols.setdefault(h, []).append(i)

    def col_values(idx):
        return [r[idx] for r in rows[1:] if r[idx] is not None]

    comuna_idx, region_idx = cols["COMUNA"][0], cols["REGION"][0]
    comuna_region = [
        [r[comuna_idx], r[region_idx]] for r in rows[1:] if r[comuna_idx] is not None
    ]

    return {
        "RESPONSABLE": col_values(cols["RESPONSABLE"][0]),
        "CANAL_COTIZACION": col_values(cols["CANAL_COTIZACION"][0]),
        "TIPO_SERVICIO": col_values(cols["TIPO_SERVICIO"][0]),
        "COMUNA_REGION": comuna_region,
        "UNIDAD": col_values(cols["UNIDAD"][0]),
        "TIPO_NEGOCIO": col_values(cols["TIPO_NEGOCIO"][0]),
    }


def build_destinatarios(destinatarios):
    cols = [
        "NOMBRE_FANTASIA", "SISTEMA_DECLARACION", "CODIGO_ESTABLECIMIENTO",
        "RAZON_SOCIAL", "RUT", "NOMBRE DE ESTABLECIMIENTO",
        "COMUNA DE ESTABLECIMIENTO", "REGION DE ESTABLECIMIETNO",
    ]
    rows = [[row.get(c) for c in cols] for row in destinatarios]
    return {"cols": cols, "rows": rows}


def main():
    if len(sys.argv) < 2:
        print("Uso: python scripts/extract_data.py ruta/al/archivo.xlsx")
        sys.exit(1)

    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)

    residuos = sheet_to_json(wb, "RESIDUOS")
    destinatarios = sheet_to_json(wb, "DESTINATARIOS")

    residuos_tree = build_residuos_tree(residuos)
    listas = build_listas(wb)
    destinatarios_compact = build_destinatarios(destinatarios)

    out = (
        "/**\n"
        " * Datos base de la plataforma, generados desde BBDD_cotizaciones_v1.xlsx\n"
        " * con scripts/extract_data.py. Si tus hojas RESIDUOS, LISTAS o DESTINATARIOS\n"
        " * cambian, vuelve a correr ese script para regenerar este archivo — no lo\n"
        " * edites a mano.\n"
        " */\n\n"
        f"export const RESIDUOS_TREE = {json.dumps(residuos_tree, ensure_ascii=False, separators=(',', ':'))};\n\n"
        f"export const LISTAS = {json.dumps(listas, ensure_ascii=False, separators=(',', ':'))};\n\n"
        f"export const DESTINATARIOS_SEED = {json.dumps(destinatarios_compact, ensure_ascii=False, separators=(',', ':'))};\n"
    )

    with open("src/data.js", "w", encoding="utf-8") as f:
        f.write(out)

    print(f"src/data.js actualizado: {len(residuos)} residuos, {len(destinatarios)} destinatarios.")


if __name__ == "__main__":
    main()
